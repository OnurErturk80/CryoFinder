from __future__ import annotations

import argparse
import posixpath
import sys

from openpyxl.utils.cell import get_column_letter, range_boundaries

from . import onedrive
from .auth import AuthError, TokenProvider
from .backup import Backuper, BackupError
from .changelog import ChangeLog
from .config import ConfigError, load_config
from .editor import parse_change, run_changes
from .excel_api import ExcelApiBackend, probe
from .graph import GraphClient, GraphError
from .memory import MemoryBackend
from .onedrive import ConflictError


def _confirm(question: str) -> bool:
    return input(f"{question} [e/H] ").strip().lower() in ("e", "evet")


def _require_tty() -> None:
    if not sys.stdin.isatty():
        raise SystemExit("Onay gerektiren işlemler etkileşimli terminalde çalışır (onay atlanamaz).")


def _guard_write(path: str, production: bool) -> None:
    """Gerçek dosya yalnızca --production ve dosya adı yazılarak onaylanırsa değiştirilebilir."""
    name = posixpath.basename(path)
    if "TEST" not in name.upper() and not production:
        raise SystemExit(f"'{name}' test dosyası değil; yazmak için --production gerekir.")
    if production:
        print(f"\n!!! GERÇEK DOSYA: {path}")
        if input("Devam etmek için dosya adını aynen yazın: ").strip() != name:
            raise SystemExit("Dosya adı eşleşmedi; iptal.")


def _print_grid(cells: dict, rng: str) -> None:
    c0, r0, c1, r1 = range_boundaries(rng)
    widths = {c: max([len(get_column_letter(c))] + [len(str(cells.get((r, c), ""))) for r in range(r0, r1 + 1)])
              for c in range(c0, c1 + 1)}
    widths = {c: min(w, 22) for c, w in widths.items()}
    print("     " + " ".join(get_column_letter(c).ljust(widths[c]) for c in range(c0, c1 + 1)))
    for r in range(r0, r1 + 1):
        row = [str(cells.get((r, c), ""))[:22].ljust(widths[c]) for c in range(c0, c1 + 1)]
        if any(x.strip() for x in row):
            print(f"{r:>4} " + " ".join(row))


def _open_backend(args, client, path):
    """auto: Excel API çalışıyorsa onu, çalışmıyorsa bellek-içi yöntemi seç."""
    meta = onedrive.get_item_by_path(client, path)
    mode = args.mode
    if mode in ("auto", "excel-api"):
        ok, steps = probe(client, meta["id"])
        if ok:
            be = ExcelApiBackend(client, meta["id"], persist=True)
            be.open()
            return meta, be, None
        if mode == "excel-api":
            raise SystemExit("Excel API çalışmıyor:\n" + "\n".join(f"  {'✓' if o else '✗'} {m}" for o, m in steps))
        print("ℹ Excel API bu dosyada çalışmıyor → bellek-içi (eTag/If-Match) yöntemi kullanılıyor.")
    meta, data = onedrive.download_consistent(client, path)
    return meta, MemoryBackend(client, meta, data), data


def cmd_login(args, cfg, client, tp):
    client.request("GET", "/me/drive", params={"$select": "id,driveType,owner"})
    print(f"✓ Giriş başarılı: {tp.username or '?'}")


def cmd_probe(args, cfg, client, tp):
    path = cfg.target_path(args.production)
    meta = onedrive.get_item_by_path(client, path)
    print(f"Dosya: {meta['name']}  ({meta['size']} bayt, eTag {meta['eTag']})")
    ok, steps = probe(client, meta["id"])
    for good, msg in steps:
        print(f"  {'✓' if good else '✗'} {msg}")
    print("\nSonuç:", "Excel API OKUMA/OTURUM çalışıyor." if ok else
          "Excel API çalışmıyor → bellek-içi eTag/If-Match yöntemi kullanılacak.")
    if ok and args.write_probe:
        _require_tty()
        _guard_write(path, args.production)
        be = ExcelApiBackend(client, meta["id"], persist=True)
        names = be.sheet_names()
        cells = be.dump(names[0])
        if not cells:
            print("Yazma testi için dolu hücre bulunamadı."); return
        (r, c), val = next(iter(sorted(cells.items())))
        addr = f"{get_column_letter(c)}{r}"
        print(f"\nYazma testi: {names[0]}!{addr} hücresine AYNI değer ({val!r}) geri yazılacak (içerik değişmez).")
        if not _confirm("Yedek alınıp yazma testi yapılsın mı?"):
            return
        log = ChangeLog(cfg.changelog_path, {"file": path, "item_id": meta["id"], "user": tp.username})
        bk = Backuper(client, cfg.backup_folder, meta["name"], lambda: onedrive.download(client, meta["id"]))
        info = bk.ensure(); log.record("backup", **info); print(f"✓ Yedek: {info['path']}")
        be.open()
        try:
            be.set_cell(names[0], addr, val if isinstance(val, str) else str(val))
            back, _ = be.get_cell(names[0], addr)
        except GraphError as e:
            log.record("failed", stage="write-probe", error=str(e))
            print(f"✗ Excel API yazması başarısız: {e}\n→ Bellek-içi yöntemi kullanın: --mode memory"); return
        finally:
            be.close()
        log.record("write-probe", sheet=names[0], cell=addr, value=val)
        print("✓ Excel API yazması çalışıyor." if str(back) == str(val) else "⚠ Yazıldı ama geri okunan değer farklı.")


def _with_backend(args, cfg, client, tp, need_write: bool):
    path = cfg.target_path(args.production)
    if need_write:
        _require_tty()
        _guard_write(path, args.production)
    meta, be, data = _open_backend(args, client, path)
    return path, meta, be, data


def cmd_sheets(args, cfg, client, tp):
    path, meta, be, _ = _with_backend(args, cfg, client, tp, False)
    try:
        print(f"{meta['name']} [{be.name}]"); [print("  -", n) for n in be.sheet_names()]
    finally:
        be.close()


def cmd_show(args, cfg, client, tp):
    path, meta, be, _ = _with_backend(args, cfg, client, tp, False)
    try:
        _print_grid(be.dump(args.sheet), args.range)
    finally:
        be.close()


def cmd_find(args, cfg, client, tp):
    path, meta, be, _ = _with_backend(args, cfg, client, tp, False)
    try:
        needle, hits = args.text.lower(), 0
        for name in be.sheet_names():
            for (r, c), v in sorted(be.dump(name).items()):
                if needle in str(v).lower():
                    print(f"{name}!{get_column_letter(c)}{r}: {v}"); hits += 1
        print(f"{hits} eşleşme.")
    finally:
        be.close()


def cmd_set(args, cfg, client, tp):
    changes = [parse_change(s) for s in args.change]
    path, meta, be, data = _with_backend(args, cfg, client, tp, True)
    log = ChangeLog(cfg.changelog_path, {"file": path, "item_id": meta["id"], "mode": be.name, "user": tp.username})
    getter = (lambda: data) if data is not None else (lambda: onedrive.download(client, meta["id"]))
    backuper = Backuper(client, cfg.backup_folder, meta["name"], getter)
    try:
        done = run_changes(be, changes, confirm=_confirm, log=log, ensure_backup=backuper.ensure)
        if not done:
            print("\nUygulanan değişiklik yok."); return
        if not be.needs_commit:
            print(f"\n✓ {len(done)} değişiklik Excel API ile uygulandı."); return
        new_bytes = be.serialize()
        problems, lost = be.verify(new_bytes)
        if problems:
            log.record("aborted", reason="unexpected-diff", details=problems)
            raise SystemExit("İPTAL (yüklenmedi): " + "; ".join(problems))
        if lost and not args.accept_feature_loss:
            log.record("aborted", reason="feature-loss", details=lost)
            raise SystemExit("İPTAL (yüklenmedi): yeniden yazım şu parçaları kaybedecek: " + ", ".join(lost) +
                             "\nBu dosya için --mode excel-api deneyin veya --accept-feature-loss ile bilerek kabul edin.")
        print(f"\n{len(done)} değişiklik OneDrive'a yüklenecek (eTag {meta['eTag']} ile kontrollü).")
        if not _confirm("Yüklensin mi?"):
            log.record("declined-upload", changes=len(done)); print("Yüklenmedi."); return
        try:
            item = be.commit(new_bytes)
        except ConflictError as e:
            log.record("conflict", error=str(e)); raise SystemExit(f"✗ {e}")
        log.record("committed", changes=len(done), new_etag=item.get("eTag"), sha256=onedrive.sha256(new_bytes),
                   backup=backuper.result["path"] if backuper.result else None)
        print("✓ Yüklendi ve doğrulandı.")
    finally:
        be.close()


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="tank_haritasi", description="OneDrive tank haritası düzenleyici (silme yok).")
    p.add_argument("--production", action="store_true", help="TEST yerine gerçek dosyayı kullan")
    p.add_argument("--mode", choices=["auto", "excel-api", "memory"], default="auto")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("login", help="Tarayıcıda giriş yap").set_defaults(fn=cmd_login)
    s = sub.add_parser("probe", help="Excel API bu dosyada çalışıyor mu?")
    s.add_argument("--write-probe", action="store_true", help="Aynı değeri geri yazarak yazmayı da dene (onaylı)")
    s.set_defaults(fn=cmd_probe)
    sub.add_parser("sheets", help="Sayfaları listele").set_defaults(fn=cmd_sheets)
    s = sub.add_parser("show", help="Aralığı göster"); s.add_argument("sheet")
    s.add_argument("--range", default="A1:L30"); s.set_defaults(fn=cmd_show)
    s = sub.add_parser("find", help="Metin ara"); s.add_argument("text"); s.set_defaults(fn=cmd_find)
    s = sub.add_parser("set", help="Hücre değiştir (her biri onaylı)")
    s.add_argument("-c", "--change", action="append", required=True, metavar="SAYFA!HÜCRE=DEĞER")
    s.add_argument("--accept-feature-loss", action="store_true")
    s.set_defaults(fn=cmd_set)
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    try:
        cfg = load_config()
        tp = TokenProvider(cfg)
        client = GraphClient(tp.token)
        args.fn(args, cfg, client, tp)
    except (ConfigError, AuthError, GraphError, BackupError, ConflictError, KeyError, ValueError) as e:
        print(f"Hata: {e}", file=sys.stderr)
        return 1
    return 0
