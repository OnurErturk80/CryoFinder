"""Yedek yöntem: dosyayı yalnızca bellekte (BytesIO) aç, düzenle, eTag/If-Match ile geri yükle.

Diske hiçbir .xlsx yazılmaz. openpyxl yeniden yazımı bazı özellikleri (grafik, resim, pivot...)
kaybedebileceği için yükleme öncesi iki doğrulama yapılır: (1) yalnızca onaylanan hücreler değişti mi,
(2) özgün dosyadaki bu tür parçalar korunuyor mu.
"""
from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook

from . import onedrive
from .sheetutil import grid_to_cells

RISKY_PREFIXES = ("xl/charts/", "xl/drawings/", "xl/media/", "xl/pivotTables/", "xl/pivotCache/",
                  "xl/slicers/", "xl/slicerCaches/", "xl/embeddings/", "xl/ctrlProps/", "xl/activeX/",
                  "xl/threadedComments/", "xl/persons/", "xl/externalLinks/")


class MemoryBackend:
    name = "memory"
    needs_commit = True

    def __init__(self, client, meta: dict, data: bytes):
        self.c, self.meta, self.original = client, meta, data
        self.wb = load_workbook(io.BytesIO(data))
        self.staged: dict[tuple[str, str], object] = {}

    def sheet_names(self) -> list[str]:
        return self.wb.sheetnames

    def _ws(self, sheet: str):
        for n in self.wb.sheetnames:
            if n.lower() == sheet.lower():
                return self.wb[n]
        raise KeyError(f"Sayfa yok: {sheet}")

    def dump(self, sheet: str, text: bool = False):
        ws = self._ws(sheet)
        fmt = (lambda v: v.date().isoformat() if hasattr(v, "date") else v) if text else (lambda v: v)
        return {(c.row, c.column): fmt(c.value) for row in ws.iter_rows() for c in row if c.value not in ("", None)}

    def get_cell(self, sheet: str, cell: str):
        c = self._ws(sheet)[cell]
        return c.value, c.data_type == "f"

    def set_cell(self, sheet: str, cell: str, value: str) -> None:
        ws = self._ws(sheet)
        ws[cell].value = value if value != "" else None
        self.staged[(ws.title, cell)] = value

    def serialize(self) -> bytes:
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()

    def verify(self, new_bytes: bytes) -> tuple[list[str], list[str]]:
        """(beklenmeyen farklar, kaybolan riskli parçalar)"""
        old, new = load_workbook(io.BytesIO(self.original)), load_workbook(io.BytesIO(new_bytes))
        problems: list[str] = []
        if old.sheetnames != new.sheetnames:
            problems.append("Sayfa listesi değişti")
        allowed = {(s, c.upper()) for s, c in self.staged}
        for name in old.sheetnames:
            if name not in new.sheetnames:
                continue
            a, b = old[name], new[name]
            coords = {(c.row, c.column) for row in a.iter_rows() for c in row if c.value is not None}
            coords |= {(c.row, c.column) for row in b.iter_rows() for c in row if c.value is not None}
            for r, col in coords:
                if a.cell(r, col).value != b.cell(r, col).value:
                    addr = b.cell(r, col).coordinate
                    if (name, addr) not in allowed:
                        problems.append(f"{name}!{addr} onaylanmadığı hâlde değişti")
        with zipfile.ZipFile(io.BytesIO(self.original)) as zo, zipfile.ZipFile(io.BytesIO(new_bytes)) as zn:
            lost = sorted(n for n in set(zo.namelist()) - set(zn.namelist()) if n.startswith(RISKY_PREFIXES))
        return problems, lost

    def commit(self, new_bytes: bytes) -> dict:
        """If-Match ile yükle, sonra geri indirip içeriği doğrula."""
        item = onedrive.upload_replace(self.c, self.meta["id"], new_bytes, self.meta["eTag"])
        back = onedrive.download(self.c, self.meta["id"])
        if onedrive.sha256(back) != onedrive.sha256(new_bytes):
            raise RuntimeError("Yükleme sonrası doğrulama başarısız: OneDrive'daki içerik gönderilenle aynı değil. "
                               "Yedek klasöründeki dosyaya bakın.")
        return item

    def open(self):  # arayüz uyumu
        pass

    def close(self):
        pass
