"""Tank haritası modeli ve boş yer önerisi.

Standart düzen: her canister 7 sütunluk bir blok (NO | SOYAD | AD | EŞİ | TARİH | HÜCRE | VİAL). NO sütununda önceden
basılı goblet konumu etiketleri (1A, 1A, 1A, 1A, 2A, ... veya 1, 1, 1, 1, ...) bulunur; aynı etiketin ardışık 4 satırı
bir konumdur (= 4 straw yeri). Bir konumda straw renkleri (MAVİ, SARI, YEŞİL, TURUNCU) birbirinden farklı olmalıdır.
Bu modül yalnızca okur ve öneri üretir; yazma işlemi servis katmanındaki onaylı akıştan geçer.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from openpyxl.utils.cell import get_column_letter

_TR = str.maketrans({"İ": "I", "ı": "I", "Ş": "S", "ş": "S", "Ü": "U", "ü": "U",
                     "Ö": "O", "ö": "O", "Ç": "C", "ç": "C", "Ğ": "G", "ğ": "G"})
POS_RE = re.compile(r"^(\d{1,2})([A-Z]?)$")
TANK_RE = re.compile(r"^TANK\s*(\d+)$")
CAN_RE = re.compile(r"^CAN+ISTER\s*(\d+)")
SLOTS = 4
COLORS = ["MAVİ", "SARI", "YEŞİL", "TURUNCU"]          # öneri sırası ve yazım
_COLOR_FOLDED = {c.translate(_TR).upper(): c for c in COLORS}
STRAW_TYPES = ["VİTRİFİT", "CRYOLOCK", "CRYOTOP"]


def fold(v) -> str:
    return re.sub(r"\s+", " ", str(v).strip().translate(_TR).upper())


def tr_upper(s: str) -> str:
    """Türkçe büyük harf: i→İ, ı→I (haritadaki isimler büyük harfle yazılıyor)."""
    return s.strip().replace("i", "İ").replace("ı", "I").upper()


def vial_text(straw_type: str, color: str) -> str:
    """Kullanıcının yazdığı biçim: '1 CRYOLOCKSARI', '1 VİTRİFİTMAVİ'."""
    return f"1 {straw_type}{color}"


def vial_kind(text) -> str | None:
    f = fold(text)
    for key, name in (("RAPIDI", "RAPIDI"), ("VITRIFIT", "VİTRİFİT"), ("CRYOLOCK", "CRYOLOCK"), ("CRYOTOP", "CRYOTOP")):
        if key in f:
            return name
    return None


def vial_color(text) -> str | None:
    f = fold(text)
    for key, color in _COLOR_FOLDED.items():
        if key in f:
            return color
    return None


@dataclass
class Row:
    row: int
    cells: dict           # alan -> (satır, sütun)
    occupied: bool
    color: str | None
    vial: str
    date_text: str = ""


@dataclass
class Position:
    sheet: str
    tank: int | None
    canister: int | None
    label: str            # "2A" ya da "2"
    number: int
    suffix: str
    rows: list[Row] = field(default_factory=list)

    @property
    def kat(self) -> str:
        """Kat sayfadan belli olur: adında 'ÜST' geçen sayfa üst kat, diğerleri alt kat.
        'A' eki yalnızca o gobletin üst katının kullanıldığını gösteren bir işarettir."""
        return "üst" if "UST" in fold(self.sheet) else "alt"

    @property
    def standard(self) -> bool:
        """Standart konum: tam 4 satır. Daha kısa/uzun gruplar (not, ek kayıt, bozuk blok) öneride kullanılmaz."""
        return len(self.rows) == SLOTS

    @property
    def name(self) -> str:
        return f"Tank {self.tank or '?'} / Canister {self.canister or '?'} / {self.label} ({self.kat} kat)"

    @property
    def free_rows(self) -> list[Row]:
        return [r for r in self.rows if not r.occupied]

    @property
    def used_colors(self) -> set[str]:
        return {r.color for r in self.rows if r.color}

    @property
    def free_colors(self) -> list[str]:
        return [c for c in COLORS if c not in self.used_colors]

    @property
    def unknown_occupied(self) -> int:
        """Dolu ama rengi okunamayan satır (rapidi renksiz olduğu için sayılmaz)."""
        return sum(1 for r in self.rows if r.occupied and not r.color and vial_kind(r.vial) != "RAPIDI")


FIELDS = ["no", "soyad", "ad", "esi", "tarih", "hucre", "vial"]


def scan_sheet(sheet: str, cells: dict[tuple[int, int], object]) -> list[Position]:
    """Bir sayfanın hücrelerinden (satır, sütun) → metin konumları çıkarır. Desteklenmeyen düzende boş döner."""
    text = {k: str(v).strip() for k, v in cells.items() if str(v).strip() != ""}
    counts = Counter(c for (r, c), v in text.items()
                     if (m := POS_RE.match(fold(v))) and 1 <= int(m[1]) <= 60)
    blocks = []
    for c, n in sorted(counts.items()):
        names = sum(1 for (r, cc), v in text.items() if cc == c + 1 and not re.search(r"\d", v))
        if n >= 20 and names >= 10:
            blocks.append(c)
    first_col = blocks[0] if blocks else 1
    tanks = []
    for (r, c), v in text.items():
        if c != first_col:          # başka sütundaki "TANK n" yazıları (küçük tank vb.) bölüm başlığı değildir
            continue
        f = fold(v)
        if m := TANK_RE.match(f):
            tanks.append((r, int(m[1])))
        elif f.startswith("KUCUK TANK"):
            tanks.append((r, 0))    # 0 = küçük tank bölgesi: bu satırlardan sonrası haritaya dahil edilmez
    tanks.sort()
    cans = sorted((r, c, int(m[1])) for (r, c), v in text.items() if (m := CAN_RE.match(fold(v))))

    def current(markers, row):
        last = None
        for mk in markers:
            if mk[0] <= row:
                last = mk
            else:
                break
        return last

    out: list[Position] = []
    for bi, c0 in enumerate(blocks):
        cols = {name: c0 + i for i, name in enumerate(FIELDS)}
        block_cans = [m for m in cans if c0 <= m[1] <= c0 + 6]
        no_rows = sorted((r, fold(v)) for (r, c), v in text.items() if c == c0 and POS_RE.match(fold(v)))
        prev = None
        for r, lab in no_rows:
            m = POS_RE.match(lab)
            tank = current(tanks, r)
            if tank and tank[1] == 0:
                continue
            can = current(block_cans, r)
            key = (tank[1] if tank else None, can[2] if can else bi + 1, lab)
            pos = out[-1] if out and prev else None
            fresh = (pos is None or prev[0] + 1 != r or prev[1] != key or len(pos.rows) >= SLOTS or pos.sheet != sheet)
            if fresh:
                pos = Position(sheet, key[0], key[1], lab, int(m[1]), m[2])
                out.append(pos)
            vial = text.get((r, cols["vial"]), "")
            occupied = any(text.get((r, cols[f]), "") for f in ("soyad", "ad", "hucre", "vial"))
            pos.rows.append(Row(r, {f: (r, cols[f]) for f in FIELDS}, occupied, vial_color(vial), vial,
                                text.get((r, cols["tarih"]), "")))
            prev = (r, key)
    return out


# ---------------- öneri
@dataclass
class Placement:
    position: Position
    rows: list[Row]
    colors: list[str | None]


@dataclass
class Plan:
    placements: list[Placement]

    @property
    def span(self) -> int:
        return len(self.placements)


def _contiguous(rows: list[Row], m: int) -> list[Row]:
    """Mümkünse m ardışık boş satır; yoksa ilk m boş satır."""
    for i in range(len(rows) - m + 1):
        chunk = rows[i:i + m]
        if all(chunk[j + 1].row == chunk[j].row + 1 for j in range(m - 1)):
            return chunk
    return rows[:m]


def suggest(positions: list[Position], n: int, colored: bool = True, limit: int = 10) -> list[Plan]:
    """n straw için en az konuma yayılan, komşu konumlardaki planlar. Aynı konumda renkler farklıdır."""
    groups: dict[tuple, list[Position]] = defaultdict(list)
    for p in positions:
        if p.standard:
            groups[(p.sheet, p.tank, p.canister, p.kat)].append(p)
    plans: list[tuple[tuple, Plan]] = []
    # Önce tüm alt katlar (canister sırasıyla), sonra üst katlar: üst kata, alt kat dolunca geçilir.
    order = sorted(groups.items(), key=lambda kv: (1 if kv[0][3] == "üst" else 0, str(kv[0][0]),
                                                   kv[0][1] or 0, kv[0][2] or 0))
    for gi, (gkey, plist) in enumerate(order):
        plist.sort(key=lambda p: p.number)
        for i, start in enumerate(plist):
            remaining, placements, used_in_plan, prev = n, [], set(), None
            for pos in plist[i:]:
                if prev is not None and pos.number != prev.number + 1:
                    break
                free = pos.free_rows
                avail = pos.free_colors if colored else [None] * len(free)
                m = min(remaining, len(free), len(avail))
                if m == 0:
                    break
                rows = _contiguous(free, m)
                if colored:
                    prefer = [c for c in avail if c not in used_in_plan] + [c for c in avail if c in used_in_plan]
                    cols = sorted(prefer[:m], key=COLORS.index)
                    used_in_plan.update(cols)
                else:
                    cols = [None] * m
                placements.append(Placement(pos, rows, cols))
                remaining -= m
                prev = pos
                if remaining == 0:
                    break
            if remaining == 0 and placements:
                plans.append(((len(placements), gi, start.number), Plan(placements)))
    plans.sort(key=lambda t: t[0])
    return [p for _, p in plans[:limit]]


def summarize(positions: list[Position]) -> list[str]:
    """Kişisel veri içermeyen özet: sayfa/tank/canister/kat başına konum sayıları."""
    by = defaultdict(list)
    for p in positions:
        by[(p.sheet, p.tank, p.canister, p.kat)].append(p)
    lines = []
    for (sheet, tank, can, kat), allp in sorted(by.items(), key=lambda kv: (kv[0][0], kv[0][1] or 0, kv[0][2] or 0, kv[0][3])):
        ps = [p for p in allp if p.standard]
        empty = sum(1 for p in ps if len(p.free_rows) == len(p.rows))
        partial = sum(1 for p in ps if 0 < len(p.free_rows) < len(p.rows))
        full = sum(1 for p in ps if not p.free_rows)
        odd = len(allp) - len(ps)
        unknown = sum(p.unknown_occupied for p in ps)
        lines.append(f"[{sheet}] Tank {tank or '?'} / Canister {can or '?'} / {kat} kat: {len(ps)} standart konum "
                     f"(tamamen boş {empty}, kısmen dolu {partial}, dolu {full}"
                     f"{', standart dışı ' + str(odd) if odd else ''}"
                     f"{', rengi okunamayan dolu satır ' + str(unknown) if unknown else ''})")
    return lines


def describe_plan(plan: Plan) -> str:
    parts = []
    for pl in plan.placements:
        cols = ", ".join(c or "renksiz" for c in pl.colors)
        rows = f"satır {pl.rows[0].row}" + (f"–{pl.rows[-1].row}" if len(pl.rows) > 1 else "")
        parts.append(f"{pl.position.name} ({rows}): {cols}   [boş renkler: {', '.join(pl.position.free_colors) or '-'}]")
    return " → ".join(parts) if len(parts) == 1 else "\n      ".join(parts)


def describe_layout(sheet: str, cells: dict, positions: list[Position]) -> list[str]:
    """Kişisel veri içermez: TANK/CANISTER işaret hücreleri ve her grubun sütun/satır/goblet aralığı."""
    lines = [f"--- {sheet} ---"]
    marks = []
    for (r, c), v in sorted(cells.items()):
        f = fold(v)
        if TANK_RE.match(f) or CAN_RE.match(f) or f.startswith("KUCUK TANK"):
            marks.append(f"{get_column_letter(c)}{r}='{str(v).strip()}'")
    lines.append("İşaret hücreleri: " + (", ".join(marks) or "yok"))
    groups = defaultdict(list)
    for p in positions:
        groups[(p.tank, p.canister, p.suffix, p.rows[0].cells["no"][1])].append(p)
    for (tank, can, suf, col), ps in sorted(groups.items(), key=lambda kv: (kv[0][3], kv[1][0].rows[0].row)):
        rows = [r.row for p in ps for r in p.rows]
        nums = sorted({p.number for p in ps})
        sizes = Counter(len(p.rows) for p in ps)
        lines.append(f"  sütun {get_column_letter(col)}: Tank {tank or '?'} / Canister {can or '?'} / "
                     f"{'A etiketli' if suf else 'harfsiz'}: satır {min(rows)}–{max(rows)}, goblet {nums[0]}–{nums[-1]} "
                     f"({len(ps)} konum; satır sayıları {dict(sorted(sizes.items()))})")
    odd = [f"{get_column_letter(p.rows[0].cells['no'][1])}{p.rows[0].row}:{p.label}({len(p.rows)} satır)"
           for p in positions if not p.standard]
    lines.append("Standart dışı konumlar: " + (", ".join(odd[:60]) + (" …" if len(odd) > 60 else "") if odd else "yok"))
    from .profile import shape
    shapes = Counter(shape(r.vial) if r.vial else "(boş)" for p in positions if p.standard
                     for r in p.rows if r.occupied and not r.color and vial_kind(r.vial) != "RAPIDI")
    if shapes:
        lines.append("Rengi okunamayan VİAL biçimleri: " + ", ".join(f"{k}×{n}" for k, n in shapes.most_common(10)))
    return lines
