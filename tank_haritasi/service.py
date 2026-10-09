"""Arayüzün kullandığı iş mantığı. CLI ile aynı güvenlik kuralları: onay, yedek, kayıt, silme yok."""
from __future__ import annotations

import re
import secrets
import threading
import time
from datetime import date

from openpyxl.utils.cell import get_column_letter

from .backup import BackupError
from .editor import Change, apply_change, evaluate, log_evaluation
from .harita import COLORS, STRAW_TYPES, Position, fold, scan_sheet, suggest, tr_upper, vial_kind, vial_text
from .sheetutil import CELL_RE

CACHE_TTL = 20.0


class EditService:
    def __init__(self, backend, backuper, log, file_label: str, production: bool = False):
        self.backend, self.backuper, self.log = backend, backuper, log
        self.file_label, self.production = file_label, production
        self.lock = threading.RLock()
        self.pending: dict[str, tuple[Change, object]] = {}
        self.staged: list[Change] = []
        self._cache: dict[str, tuple[float, dict]] = {}
        self._backup_logged = False
        self._plans: dict[str, object] = {}
        self._regs: dict[str, list[Change]] = {}
        self._rems: dict[str, dict] = {}

    # ---- okuma
    def _cells(self, sheet: str, fresh: bool = False) -> dict:
        hit = self._cache.get(sheet)
        if hit and not fresh and time.time() - hit[0] < CACHE_TTL:
            return hit[1]
        cells = self.backend.dump(sheet, text=True)
        self._cache[sheet] = (time.time(), cells)
        return cells

    def state(self) -> dict:
        with self.lock:
            return {"file": self.file_label, "mode": self.backend.name, "production": self.production,
                    "sheets": self.backend.sheet_names(), "staged": len(self.staged),
                    "needs_commit": self.backend.needs_commit,
                    "backup": self.backuper.result["path"] if self.backuper.result else None}

    def sheet_view(self, sheet: str, row0: int = 1, nrows: int = 50, fresh: bool = False) -> dict:
        with self.lock:
            cells = self._cells(sheet, fresh)
            max_row = max((r for r, _ in cells), default=1)
            max_col = max((c for _, c in cells), default=1)
            row0 = max(1, min(row0, max_row))
            nrows = max(1, min(nrows, 5000))
            rows = [[str(cells.get((r, c), "")) for c in range(1, max_col + 1)]
                    for r in range(row0, min(row0 + nrows, max_row + 1))]
            return {"sheet": sheet, "row0": row0, "max_row": max_row, "max_col": max_col,
                    "cols": [get_column_letter(c) for c in range(1, max_col + 1)], "rows": rows}

    def search(self, query: str, limit: int = 200) -> list[dict]:
        needle = query.strip().lower()
        if len(needle) < 2:
            return []
        hits: list[dict] = []
        with self.lock:
            for name in self.backend.sheet_names():
                for (r, c), v in sorted(self._cells(name).items()):
                    if needle in str(v).lower():
                        hits.append({"sheet": name, "cell": f"{get_column_letter(c)}{r}", "row": r, "value": str(v)})
                        if len(hits) >= limit:
                            return hits
        return hits

    # ---- yazma: öner → (kullanıcı onayı) → uygula
    def propose(self, sheet: str, cell: str, value: str) -> dict:
        cell = cell.strip().upper()
        if not CELL_RE.match(cell):
            return {"ok": False, "message": "Geçersiz hücre adresi."}
        ch = Change(sheet, cell, value)
        with self.lock:
            ev = evaluate(self.backend, ch)
            if ev["kind"] != "ok":
                log_evaluation(self.log, ch, ev)
                return {"ok": False, "message": ev["message"]}
            pid = secrets.token_urlsafe(9)
            self.pending[pid] = (ch, ev["old"])
            return {"ok": True, "id": pid, "sheet": sheet, "cell": cell,
                    "old": "" if ev["old"] is None else str(ev["old"]), "new": value}

    def decline(self, pid: str) -> dict:
        with self.lock:
            item = self.pending.pop(pid, None)
            if item:
                ch, old = item
                self.log.record("declined", sheet=ch.sheet, cell=ch.cell, old=old, new=ch.new_value)
        return {"ok": True}

    def apply(self, pid: str) -> dict:
        """Kullanıcı arayüzde 'Onayla'ya bastığında çağrılır. Tek kullanımlık kimlik."""
        with self.lock:
            item = self.pending.pop(pid, None)
            if not item:
                return {"ok": False, "message": "Bu öneri artık geçerli değil; yeniden başlatın."}
            ch, old = item
            ev = evaluate(self.backend, ch)
            if ev["kind"] != "ok" or ev["old"] != old:
                return {"ok": False, "message": "Hücre siz incelerken değişmiş. Yeniden gözden geçirin."}
            try:
                info = self.backuper.ensure()
            except BackupError as e:
                return {"ok": False, "message": str(e)}
            if not self._backup_logged:
                self.log.record("backup", **info)
                self._backup_logged = True
            try:
                apply_change(self.backend, ch, old, self.log)
            except Exception as e:  # noqa: BLE001
                return {"ok": False, "message": f"Yazılamadı: {e}"}
            if self.backend.needs_commit:
                self.staged.append(ch)
            self._cache.pop(ch.sheet, None)
            return {"ok": True, "backup": info["path"], "staged": len(self.staged),
                    "message": "Hazırlandı; 'OneDrive'a yükle' ile tamamlayın." if self.backend.needs_commit
                    else "Uygulandı."}

    # ---- yalnızca bellek modu
    def prepare_commit(self) -> dict:
        with self.lock:
            if not self.backend.needs_commit or not self.staged:
                return {"ok": False, "message": "Yüklenecek değişiklik yok."}
            new_bytes = self.backend.serialize()
            problems, lost = self.backend.verify(new_bytes)
            return {"ok": not problems and not lost, "count": len(self.staged), "problems": problems,
                    "lost": lost, "etag": self.backend.meta["eTag"],
                    "changes": [f"{c.sheet}!{c.cell}" for c in self.staged]}

    def commit(self) -> dict:
        from .onedrive import ConflictError, sha256
        with self.lock:
            prep = self.prepare_commit()
            if not prep["ok"]:
                self.log.record("aborted", reason="verification", details=prep.get("problems") or prep.get("lost"))
                return {"ok": False, "message": "Doğrulama başarısız; yüklenmedi: " +
                        "; ".join((prep.get("problems") or []) + (prep.get("lost") or []))}
            new_bytes = self.backend.serialize()
            try:
                item = self.backend.commit(new_bytes)
            except ConflictError as e:
                self.log.record("conflict", error=str(e))
                return {"ok": False, "message": str(e)}
            self.log.record("committed", changes=len(self.staged), new_etag=item.get("eTag"),
                            sha256=sha256(new_bytes), backup=self.backuper.result["path"])
            self.staged.clear()
            self.backend.meta = {**self.backend.meta, "eTag": item.get("eTag", self.backend.meta["eTag"])}
            self.backend.original = new_bytes  # sonraki karşılaştırmalar yeni dosyayı esas alsın
            self._cache.clear()
            return {"ok": True, "message": "Yüklendi ve doğrulandı. Sayfayı yeniden açın."}

    def recent_log(self, path, n: int = 15) -> list[str]:
        try:
            return path.read_text(encoding="utf-8").splitlines()[-n:]
        except FileNotFoundError:
            return []


    # ---- yeni hasta kaydı: yer öner → önizle → tek onayla yaz
    def _scan_all(self) -> list[Position]:
        out: list[Position] = []
        for name in self.backend.sheet_names():
            out += scan_sheet(name, self._cells(name, fresh=True))
        return out

    def reg_options(self) -> dict:
        with self.lock:
            pos = [p for p in self._scan_all() if p.standard and p.tank]
            tanks = sorted({p.tank for p in pos})
            return {"tanks": tanks, "types": STRAW_TYPES, "colors": COLORS,
                    "free_rows": {str(t): sum(len(p.free_rows) for p in pos if p.tank == t) for t in tanks}}

    @staticmethod
    def _pkey(p: Position) -> tuple:
        return (p.sheet, p.tank, p.canister, p.label, p.rows[0].row)

    def reg_suggest(self, tank: int, n: int, kat: str) -> dict:
        if not 1 <= n <= 12:
            return {"ok": False, "message": "Straw sayısı 1–12 olmalı."}
        if kat not in ("otomatik", "alt", "ust"):
            return {"ok": False, "message": "Geçersiz kat."}
        with self.lock:
            pool = [p for p in self._scan_all() if p.standard and p.tank == tank
                    and (kat == "otomatik" or p.kat == ("üst" if kat == "ust" else "alt"))]
            plans = suggest(pool, n, colored=True, limit=8)
            self._plans = {}
            out = []
            for i, plan in enumerate(plans):
                pid = f"p{i}"
                self._plans[pid] = plan
                out.append({"id": pid, "span": plan.span, "placements": [{
                    "sheet": pl.position.sheet, "canister": pl.position.canister, "label": pl.position.label,
                    "kat": pl.position.kat, "rows": [r.row for r in pl.rows], "colors": pl.colors,
                    "free_colors": pl.position.free_colors, "key": list(self._pkey(pl.position)),
                    "free_row_count": len(pl.position.free_rows)} for pl in plan.placements]})
            return {"ok": True, "plans": out,
                    "message": "" if out else "Bu tankta, seçilen kat ve straw sayısı için uygun yer yok."}

    def _date_format(self, positions: list[Position], sheet: str):
        for p in positions:
            if p.sheet == sheet:
                for r in p.rows:
                    if r.date_text:
                        rr, cc = r.cells["tarih"]
                        try:
                            return self.backend.get_number_format(sheet, f"{get_column_letter(cc)}{rr}")
                        except Exception:  # noqa: BLE001 - biçim okunamazsa hücre varsayılanı kalır
                            return None
        return None

    def reg_preview(self, plan_id: str, straw_type: str, straws: list[dict], patient: dict) -> dict:
        with self.lock:
            plan = self._plans.get(plan_id)
            if not plan:
                return {"ok": False, "message": "Öneri süresi dolmuş; yeniden yer önerin."}
            if straw_type not in STRAW_TYPES:
                return {"ok": False, "message": "Geçersiz straw türü."}
            soyad, ad, esi = (tr_upper(str(patient.get(k, ""))) for k in ("soyad", "ad", "esi"))
            if not soyad or not ad:
                return {"ok": False, "message": "Soyad ve ad zorunlu."}
            try:
                tarih = date.fromisoformat(str(patient.get("tarih", "")))
            except ValueError:
                return {"ok": False, "message": "Geçersiz tarih."}
            for v in (soyad, ad, esi):
                if len(v) > 60 or v.startswith("=") or re.search(r"[\x00-\x1f]", v):
                    return {"ok": False, "message": "Ad/soyad geçersiz karakter içeriyor."}
            fresh = self._scan_all()
            by_key = {self._pkey(p): p for p in fresh}
            expected = sum(len(pl.rows) for pl in plan.placements)
            if len(straws) != expected:
                return {"ok": False, "message": "Straw sayısı öneriyle uyuşmuyor."}
            fmt = self._date_format(fresh, plan.placements[0].position.sheet)
            changes: list[Change] = []
            i = 0
            for pl in plan.placements:
                pos = by_key.get(self._pkey(pl.position))
                if pos is None:
                    return {"ok": False, "message": "Harita değişmiş; yeniden yer önerin."}
                free = {r.row for r in pos.free_rows}
                colors = []
                for row in pl.rows:
                    st = straws[i]; i += 1
                    color, hucre = st.get("color"), str(st.get("hucre", "")).strip()
                    if row.row not in free:
                        return {"ok": False, "message": f"{pos.name}: satır {row.row} artık boş değil. Yeniden yer önerin."}
                    if color not in COLORS or color in colors or color in pos.used_colors:
                        return {"ok": False, "message": f"{pos.name}: renk seçimi geçersiz (bu goblette aynı renk iki kez "
                                                        "olamaz veya renk zaten kullanımda)."}
                    if not hucre or len(hucre) > 80 or hucre.startswith("=") or re.search(r"[\x00-\x1f]", hucre):
                        return {"ok": False, "message": "HÜCRE bilgisi boş veya geçersiz."}
                    colors.append(color)
                    vals = {"soyad": soyad, "ad": ad, "esi": esi or None, "tarih": tarih, "hucre": hucre,
                            "vial": vial_text(straw_type, color)}
                    for f, v in vals.items():
                        if v is None:
                            continue
                        rr, cc = row.cells[f]
                        changes.append(Change(pos.sheet, f"{get_column_letter(cc)}{rr}", v,
                                              fmt if f == "tarih" else None))
            preview = []
            for ch in changes:
                ev = evaluate(self.backend, ch)
                if ev["kind"] not in ("ok", "unchanged") or str(ev.get("old") or "").strip():
                    return {"ok": False, "message": f"{ch.sheet}!{ch.cell} dolu veya yazılamaz; yeniden yer önerin."}
                shown = ch.new_value.strftime("%d.%m.%Y") if isinstance(ch.new_value, date) else ch.new_value
                preview.append({"where": f"{ch.sheet}!{ch.cell}", "old": "", "new": shown})
            rid = secrets.token_urlsafe(9)
            self._regs[rid] = changes
            return {"ok": True, "id": rid, "changes": preview, "count": len(changes)}

    def reg_apply(self, rid: str) -> dict:
        with self.lock:
            changes = self._regs.pop(rid, None)
            if not changes:
                return {"ok": False, "message": "Bu kayıt artık geçerli değil; yeniden önizleyin."}
            for ch in changes:                      # yazmadan önce hepsinin hâlâ boş olduğunu doğrula
                ev = evaluate(self.backend, ch)
                if ev["kind"] not in ("ok", "unchanged") or str(ev.get("old") or "").strip():
                    return {"ok": False, "message": "Hedef hücreler siz onaylarken değişmiş. Hiçbir şey yazılmadı; "
                                                    "yeniden yer önerin."}
            try:
                info = self.backuper.ensure()
            except BackupError as e:
                return {"ok": False, "message": str(e)}
            if not self._backup_logged:
                self.log.record("backup", **info)
                self._backup_logged = True
            done = 0
            for ch in changes:
                try:
                    apply_change(self.backend, ch, None, self.log)
                except Exception as e:  # noqa: BLE001
                    return {"ok": False, "message": f"{done}/{len(changes)} hücre yazıldıktan sonra hata: {e}. "
                                                    "Kısmi kayıt olabilir; haritayı kontrol edin (ayrıntı kayıt dosyasında)."}
                done += 1
                if self.backend.needs_commit:
                    self.staged.append(ch)
            self.log.record("registration", cells=done, backup=info["path"])
            self._cache.clear()
            return {"ok": True, "backup": info["path"], "count": done,
                    "message": "Hazırlandı; 'OneDrive'a yükle' ile tamamlayın." if self.backend.needs_commit
                    else f"Kayıt yazıldı ({done} hücre)."}


    # ---- hasta çıkarma: satırın hücrelerini boşaltır (dosya/satır silinmez)
    CLEAR_FIELDS = ("soyad", "ad", "esi", "tarih", "hucre", "vial")

    def rem_search(self, query: str) -> dict:
        q = fold(query)
        if len(q) < 2:
            return {"ok": False, "message": "En az 2 harf yazın."}
        hits = []
        with self.lock:
            for p in self._scan_all():
                for r in p.rows:
                    full = (fold(f"{r.soyad} {r.ad}"), fold(f"{r.ad} {r.soyad}"))
                    if (r.soyad or r.ad) and any(q in f for f in full):
                        hits.append({"sheet": p.sheet, "tank": p.tank, "canister": p.canister, "label": p.label,
                                     "kat": p.kat, "row": r.row, "col": r.cells["no"][1], "soyad": r.soyad, "ad": r.ad, "esi": r.esi,
                                     "tarih": r.date_text, "hucre": r.hucre, "vial": r.vial})
        hits.sort(key=lambda h: (h["sheet"], h["tank"] or 0, h["canister"] or 0, h["row"]))
        return {"ok": True, "hits": hits[:200], "truncated": len(hits) > 200}

    def rem_preview(self, items: list[dict]) -> dict:
        if not items or len(items) > 40:
            return {"ok": False, "message": "1–40 satır seçin."}
        with self.lock:
            # Aynı sayfada canister blokları yan yana olduğundan satır, NO sütunuyla (blok) birlikte tanımlanır.
            rows = {(p.sheet, r.row, r.cells["no"][1]): r for p in self._scan_all() for r in p.rows}
            changes, olds, preview = [], {}, []
            seen = set()
            for it in items:
                key = (str(it.get("sheet")), int(it.get("row")), int(it.get("col")))
                row = rows.get(key)
                if row is None or not (row.soyad or row.ad) or key in seen:
                    return {"ok": False, "message": "Satır bulunamadı veya artık dolu değil; aramayı yenileyin."}
                seen.add(key)
                shown = {"soyad": row.soyad, "ad": row.ad, "esi": row.esi, "tarih": row.date_text,
                         "hucre": row.hucre, "vial": row.vial}
                for f in self.CLEAR_FIELDS:
                    rr, cc = row.cells[f]
                    ch = Change(key[0], f"{get_column_letter(cc)}{rr}", "")
                    ev = evaluate(self.backend, ch)
                    if ev["kind"] == "unchanged":
                        continue
                    if ev["kind"] != "ok":
                        return {"ok": False, "message": f"{ch.sheet}!{ch.cell}: {ev['message']}"}
                    changes.append(ch)
                    olds[(ch.sheet, ch.cell)] = ev["old"]
                    preview.append({"where": f"{ch.sheet}!{ch.cell}", "old": shown[f], "new": ""})
            if not changes:
                return {"ok": False, "message": "Temizlenecek dolu hücre yok."}
            rid = secrets.token_urlsafe(9)
            self._rems[rid] = {"changes": changes, "olds": olds}
            return {"ok": True, "id": rid, "changes": preview, "count": len(changes), "rows": len(seen)}

    def rem_apply(self, rid: str) -> dict:
        with self.lock:
            rem = self._rems.pop(rid, None)
            if not rem:
                return {"ok": False, "message": "Bu işlem artık geçerli değil; yeniden önizleyin."}
            for ch in rem["changes"]:                # onay sırasında hücreler değişmişse hiçbir şey silinmez
                ev = evaluate(self.backend, ch)
                if ev["kind"] != "ok" or ev["old"] != rem["olds"][(ch.sheet, ch.cell)]:
                    return {"ok": False, "message": "Satırlar siz onaylarken değişmiş. Hiçbir şey silinmedi; "
                                                    "aramayı yenileyip tekrar deneyin."}
            try:
                info = self.backuper.ensure()
            except BackupError as e:
                return {"ok": False, "message": str(e)}
            if not self._backup_logged:
                self.log.record("backup", **info)
                self._backup_logged = True
            done = 0
            for ch in rem["changes"]:
                try:
                    apply_change(self.backend, ch, rem["olds"][(ch.sheet, ch.cell)], self.log)
                except Exception as e:  # noqa: BLE001
                    return {"ok": False, "message": f"{done}/{len(rem['changes'])} hücre temizlendikten sonra hata: {e}. "
                                                    "Kısmi temizlik olabilir; haritayı kontrol edin."}
                done += 1
                if self.backend.needs_commit:
                    self.staged.append(ch)
            self.log.record("removal", cells=done, backup=info["path"])
            self._cache.clear()
            return {"ok": True, "count": done, "backup": info["path"],
                    "message": "Hazırlandı; 'OneDrive'a yükle' ile tamamlayın." if self.backend.needs_commit
                    else f"{done} hücre temizlendi. Yedek: {info['path']}"}


    # ---- görsel tank haritası (adlar harita verisinde yer almaz; yalnızca tıklama/aramada döner)
    @staticmethod
    def _slot(p: Position, r) -> dict:
        if not r.occupied:
            state = "empty"
        elif r.color:
            state = "color"
        elif vial_kind(r.vial) == "RAPIDI":
            state = "rapidi"
        else:
            state = "other"
        return {"id": f"{r.row}:{r.cells['no'][1]}:{p.sheet}", "row": r.row, "state": state, "color": r.color,
                "vial": r.vial}

    def map_view(self, tank: int) -> dict:
        with self.lock:
            allp = [p for p in self._scan_all() if p.tank == tank]
            pos = [p for p in allp if p.standard]
            if not pos:
                return {"ok": False, "message": "Bu tank için standart konum bulunamadı."}
            cans: dict[int, dict[int, dict]] = {}
            for p in sorted(pos, key=lambda p: (p.canister or 0, p.number)):
                g = cans.setdefault(p.canister or 0, {}).setdefault(p.number, {"n": p.number, "alt": None, "ust": None})
                key = "ust" if p.kat == "üst" else "alt"
                if g[key] is None:
                    g[key] = {"label": p.label, "slots": [self._slot(p, r) for r in p.rows]}
            stats = {"slots": 0, "free": 0, "rapidi": 0, "other": 0, **{c: 0 for c in COLORS}}
            for p in pos:
                for r in p.rows:
                    stats["slots"] += 1
                    st = self._slot(p, r)
                    if st["state"] == "empty":
                        stats["free"] += 1
                    elif st["state"] == "color":
                        stats[st["color"]] += 1
                    else:
                        stats[st["state"]] += 1
            return {"ok": True, "tank": tank, "has_ust": any(p.kat == "üst" for p in pos), "stats": stats,
                    "nonstandard": len(allp) - len(pos),
                    "canisters": [{"canister": c, "goblets": [g for _, g in sorted(gs.items())]}
                                  for c, gs in sorted(cans.items())]}

    def _row_by_id(self, slot_id: str):
        try:
            row, col, sheet = slot_id.split(":", 2)
            row, col = int(row), int(col)
        except ValueError:
            return None
        for p in self._scan_all():
            for r in p.rows:
                if p.sheet == sheet and r.row == row and r.cells["no"][1] == col:
                    return p, r
        return None

    def map_slot(self, slot_id: str) -> dict:
        with self.lock:
            hit = self._row_by_id(slot_id)
            if not hit:
                return {"ok": False, "message": "Satır bulunamadı; haritayı yenileyin."}
            p, r = hit
            st = self._slot(p, r)
            return {"ok": True, "where": f"{p.name} · satır {r.row}", "state": st["state"], "color": r.color,
                    "vial": r.vial, "soyad": r.soyad, "ad": r.ad, "esi": r.esi, "tarih": r.date_text,
                    "hucre": r.hucre}

    def map_find(self, query: str, tank: int | None = None) -> dict:
        q = fold(query)
        if len(q) < 2:
            return {"ok": False, "message": "En az 2 harf yazın."}
        hits = []
        with self.lock:
            for p in self._scan_all():
                if tank is not None and p.tank != tank:
                    continue
                for r in p.rows:
                    if (r.soyad or r.ad) and any(q in fold(f) for f in (f"{r.soyad} {r.ad}", f"{r.ad} {r.soyad}")):
                        hits.append({"id": f"{r.row}:{r.cells['no'][1]}:{p.sheet}",
                                     "where": f"Tank {p.tank or '?'} / Canister {p.canister or '?'} / {p.label} "
                                              f"({p.kat} kat) · satır {r.row}", "vial": r.vial})
        return {"ok": True, "hits": hits[:100]}
