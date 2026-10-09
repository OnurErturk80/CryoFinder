import io
import re
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tank_haritasi import onedrive
from tank_haritasi.backup import Backuper, BackupError
from tank_haritasi.changelog import ChangeLog
from tank_haritasi.editor import parse_change, run_changes
from tank_haritasi.excel_api import ExcelApiBackend, probe
from tank_haritasi.graph import ForbiddenOperation, GraphClient
from tank_haritasi.memory import MemoryBackend
from tests.fake_onedrive import FakeOneDrive, make_xlsx

PATH = "TANK_HARITASI_TEST.xlsx"
SEED = {"A1": "Kanister 1", "B1": "Boş", "C3": "=1+1"}


def setup(excel_api=False, files=None):
    fake = FakeOneDrive(files or {PATH: make_xlsx(SEED)}, excel_api=excel_api)
    return fake, GraphClient(lambda: "tok", session=fake, sleep=lambda s: None)


def logger(tmp_path):
    return ChangeLog(tmp_path / "log.jsonl", {"file": PATH})


def test_no_delete_anywhere():
    _, c = setup()
    with pytest.raises(ForbiddenOperation):
        c.request("DELETE", "/me/drive/items/ID1")
    src = "\n".join(p.read_text(encoding="utf-8") for p in Path("tank_haritasi").glob("*.py"))
    methods = set(re.findall(r"""\.request\(\s*["'](\w+)["']""", src))
    assert methods and methods <= {"GET", "POST", "PUT", "PATCH"}
    assert not re.search(r"\.delete\(|os\.remove|unlink|rmtree|shutil\.move", src)


def test_probe_detects_missing_excel_api():
    fake, c = setup(excel_api=False)
    ok, steps = probe(c, "ID1")
    assert not ok and not steps[-1][0]


def test_probe_ok_when_supported():
    fake, c = setup(excel_api=True)
    assert probe(c, "ID1")[0]


def test_parse_change():
    ch = parse_change("Tank1!b3=Yeni = değer")
    assert (ch.sheet, ch.cell, ch.new_value) == ("Tank1", "B3", "Yeni = değer")
    with pytest.raises(ValueError):
        parse_change("B3=1")


def test_memory_flow_backup_commit_log_and_decline(tmp_path):
    fake, c = setup()
    meta, data = onedrive.download_consistent(c, PATH)
    be = MemoryBackend(c, meta, data)
    bk = Backuper(c, "Yedekler", meta["name"], lambda: data, now=lambda: datetime(2026, 10, 9, 12, 0, 0))
    answers = iter([True, False])
    changes = [parse_change("Tank1!B1=Dolu"), parse_change("Tank1!A1=Başka")]
    done = run_changes(be, changes, confirm=lambda q: next(answers), log=logger(tmp_path),
                       ensure_backup=bk.ensure, say=lambda *_: None)
    assert [d.cell for d in done] == ["B1"]
    problems, lost = be.verify(be.serialize())
    assert problems == [] and lost == []
    be.commit(be.serialize())
    wb = load_workbook(io.BytesIO(fake.items[fake.by_path(PATH)]["data"]))
    assert wb["Tank1"]["B1"].value == "Dolu" and wb["Tank1"]["A1"].value == "Kanister 1"
    backup = fake.items[fake.by_path("Yedekler/TANK_HARITASI_TEST_yedek_2026-10-09_120000.xlsx")]
    assert backup["data"] == data  # yedek, değişiklikten ÖNCEKİ içerik
    events = [l for l in (tmp_path / "log.jsonl").read_text(encoding="utf-8").splitlines()]
    assert any('"declined"' in e for e in events) and any('"staged"' in e for e in events)
    assert not any(m == "DELETE" for m, _ in fake.log)


def test_etag_conflict_writes_nothing(tmp_path):
    fake, c = setup()
    meta, data = onedrive.download_consistent(c, PATH)
    be = MemoryBackend(c, meta, data)
    be.set_cell("Tank1", "B1", "Benim")
    theirs = make_xlsx({"A1": "Başkasının düzenlemesi"})
    fake.external_edit(PATH, theirs)
    with pytest.raises(onedrive.ConflictError):
        be.commit(be.serialize())
    assert fake.items[fake.by_path(PATH)]["data"] == theirs


def test_no_write_without_approval_or_backup(tmp_path):
    fake, c = setup()
    meta, data = onedrive.download_consistent(c, PATH)
    be = MemoryBackend(c, meta, data)
    # onay yok → yedek de yazma da yok
    run_changes(be, [parse_change("Tank1!B1=X")], confirm=lambda q: False, log=logger(tmp_path),
                ensure_backup=lambda: pytest.fail("yedek istenmemeliydi"), say=lambda *_: None)
    assert be.staged == {}
    # yedek başarısız → yazma yok
    def boom(): raise BackupError("disk dolu")
    with pytest.raises(BackupError):
        run_changes(be, [parse_change("Tank1!B1=X")], confirm=lambda q: True, log=logger(tmp_path),
                    ensure_backup=boom, say=lambda *_: None)
    assert be.staged == {}


def test_formula_cells_and_values_refused(tmp_path):
    fake, c = setup()
    meta, data = onedrive.download_consistent(c, PATH)
    be = MemoryBackend(c, meta, data)
    done = run_changes(be, [parse_change("Tank1!C3=5"), parse_change("Tank1!B1==SUM(A1)")],
                       confirm=lambda q: True, log=logger(tmp_path), ensure_backup=lambda: {"path": "x"},
                       say=lambda *_: None)
    assert done == []


def test_verify_catches_unapproved_diff(tmp_path):
    fake, c = setup()
    meta, data = onedrive.download_consistent(c, PATH)
    be = MemoryBackend(c, meta, data)
    be.set_cell("Tank1", "B1", "ok")
    be.wb["Tank1"]["A1"] = "sızan değişiklik"
    problems, _ = be.verify(be.serialize())
    assert problems == ["Tank1!A1 onaylanmadığı hâlde değişti"]


def test_excel_api_mode_applies_and_logs(tmp_path):
    fake, c = setup(excel_api=True)
    be = ExcelApiBackend(c, "ID1"); be.open()
    bk = Backuper(c, "Yedekler", PATH, lambda: onedrive.download(c, "ID1"))
    done = run_changes(be, [parse_change("Tank1!B1=Excel API")], confirm=lambda q: True,
                       log=logger(tmp_path), ensure_backup=bk.ensure, say=lambda *_: None)
    be.close()
    assert len(done) == 1
    assert load_workbook(io.BytesIO(fake.items["ID1"]["data"]))["Tank1"]["B1"].value == "Excel API"
    assert fake.by_path(next(p for p in [i["path"] for i in fake.items.values()] if p.startswith("Yedekler/T")))


def test_gitignore_protects_secrets():
    text = Path(".gitignore").read_text()
    for pat in (".env", ".token_cache", "degisiklik_kaydi", "*.xlsx"):
        assert pat in text


# ---- arayüz servisi ve yerel sunucu güvenliği
import json
import urllib.error
import urllib.request

from tank_haritasi.service import EditService
from tank_haritasi.webui import make_server, serve_in_thread


def make_service(tmp_path, excel_api=True):
    fake, c = setup(excel_api=excel_api)
    if excel_api:
        be = ExcelApiBackend(c, "ID1")
        data = None
    else:
        meta, data = onedrive.download_consistent(c, PATH)
        be = MemoryBackend(c, meta, data)
    getter = (lambda: data) if data else (lambda: onedrive.download(c, "ID1"))
    bk = Backuper(c, "Yedekler", PATH, getter, now=lambda: datetime(2026, 10, 9, 12, 0, 0))
    return fake, EditService(be, bk, logger(tmp_path), PATH)


def test_service_propose_apply_backup_and_single_use(tmp_path):
    fake, svc = make_service(tmp_path)
    p = svc.propose("Tank1", "B1", "Yeni")
    assert p["ok"] and p["old"] == "Boş" and p["new"] == "Yeni"
    assert not any(v["path"].startswith("Yedekler/T") for v in fake.items.values())  # onaydan önce yedek/yazma yok
    r = svc.apply(p["id"])
    assert r["ok"] and r["backup"].startswith("Yedekler/")
    assert load_workbook(io.BytesIO(fake.items["ID1"]["data"]))["Tank1"]["B1"].value == "Yeni"
    assert not svc.apply(p["id"])["ok"]  # tek kullanımlık


def test_service_refuses_formula_and_detects_stale(tmp_path):
    fake, svc = make_service(tmp_path)
    assert not svc.propose("Tank1", "C3", "5")["ok"]
    assert not svc.propose("Tank1", "B1", "=1+1")["ok"]
    p = svc.propose("Tank1", "B1", "Yeni")
    fake.external_edit(PATH, make_xlsx({"B1": "başkası yazdı"}))
    assert not svc.apply(p["id"])["ok"]
    assert load_workbook(io.BytesIO(fake.items["ID1"]["data"]))["Tank1"]["B1"].value == "başkası yazdı"


def test_service_decline_writes_nothing(tmp_path):
    fake, svc = make_service(tmp_path)
    p = svc.propose("Tank1", "B1", "X")
    svc.decline(p["id"])
    assert not svc.apply(p["id"])["ok"]
    assert "declined" in (tmp_path / "log.jsonl").read_text(encoding="utf-8")


def test_service_memory_commit_flow(tmp_path):
    fake, svc = make_service(tmp_path, excel_api=False)
    svc.apply(svc.propose("Tank1", "B1", "Bellek")["id"])
    assert svc.state()["staged"] == 1
    assert svc.prepare_commit()["ok"]
    assert svc.commit()["ok"]
    assert load_workbook(io.BytesIO(fake.items["ID1"]["data"]))["Tank1"]["B1"].value == "Bellek"
    svc.apply(svc.propose("Tank1", "B1", "İkinci")["id"])
    assert svc.commit()["ok"]  # eTag güncellendiği için ikinci yükleme çakışmaz


def _call(url, token=None, host=None, body=None, origin=None, ctype="application/json"):
    req = urllib.request.Request(url, data=None if body is None else json.dumps(body).encode(),
                                 method="POST" if body is not None else "GET")
    if token: req.add_header("X-Session-Token", token)
    if host: req.add_header("Host", host)
    if origin: req.add_header("Origin", origin)
    if body is not None: req.add_header("Content-Type", ctype)
    try:
        with urllib.request.urlopen(req) as r:
            return r.status, r.read().decode()
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


def test_webui_security_and_happy_path(tmp_path):
    fake, svc = make_service(tmp_path)
    srv, token = make_server(svc, tmp_path / "log.jsonl")
    serve_in_thread(srv)
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        assert _call(base + "/api/state")[0] == 403                                   # anahtarsız
        assert _call(base + "/api/state", token="yanlis")[0] == 403
        assert _call(base + "/api/state", token=token, host="evil.example")[0] == 403  # DNS rebinding
        assert _call(base + "/api/propose", token=token, body={}, origin="http://evil.example")[0] == 403
        assert _call(base + "/api/propose", token=token, body={"sheet": "Tank1", "cell": "B1"}, ctype="text/plain")[0] == 415
        code, page = _call(base + "/")
        assert code == 200 and token in page
        assert json.loads(_call(base + "/api/state", token=token)[1])["mode"] == "excel-api"
        view = json.loads(_call(base + "/api/sheet?name=Tank1", token=token)[1])
        assert view["rows"][0][0] == "Kanister 1"
        p = json.loads(_call(base + "/api/propose", token=token, body={"sheet": "Tank1", "cell": "B1", "value": "Web"})[1])
        assert json.loads(_call(base + "/api/apply", token=token, body={"id": p["id"]})[1])["ok"]
        assert _call(base + "/api/delete", token=token, body={})[0] == 404             # silme ucu yok
    finally:
        srv.shutdown()


def test_profile_hides_personal_data():
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from tank_haritasi.profile import profile_workbook
    wb = Workbook(); ws = wb.active; ws.title = "TANK 1"
    ws["A1"] = "TANK 1"; ws["A3"] = "CANISTER 1"
    for i, h in enumerate(["NO", "SOYAD", "AD", "TARİH", "VİAL"], 1):
        ws.cell(5, i, h)
    for r, (no, soyad, ad, vial) in enumerate([("1A", "GİZLİSOYAD", "GİZLİAD", "RAPIDI 1"),
                                               ("1B", "BAŞKASOYAD", "BAŞKAAD", "VİTRİFİT MAVİ")], 6):
        ws.cell(r, 1, no); ws.cell(r, 2, soyad); ws.cell(r, 3, ad); ws.cell(r, 4, 44597); ws.cell(r, 5, vial)
        ws.cell(r, 5).fill = PatternFill("solid", fgColor="0000FF")
    b = io.BytesIO(); wb.save(b)
    out = "\n".join(profile_workbook(b.getvalue()))
    for secret in ("GİZLİSOYAD", "GİZLİAD", "BAŞKASOYAD", "BAŞKAAD"):
        assert secret not in out
    assert "CANISTER 1" in out and "VİTRİFİT MAVİ" in out and "0000FF" in out and "SOYAD" in out


# ---- harita modeli ve yer önerisi
from tank_haritasi.harita import COLORS, scan_sheet, suggest, summarize, vial_color, vial_text


def make_map_cells():
    """Tek sayfa, TANK 1, iki canister bloğu (sütun A ve I); her blokta 1A..8A (4'er satır)."""
    cells = {(1, 1): "TANK 1", (3, 1): "CANISTER 1", (3, 9): "CANISTER 2"}
    def put(block_col, row, soyad, vial):
        cells[(row, block_col + 1)], cells[(row, block_col + 2)] = soyad, "Ad"
        cells[(row, block_col + 5)], cells[(row, block_col + 6)] = "D5 (4AA)", vial
    for col in (1, 9):
        for g in range(1, 9):
            for k in range(4):
                cells[(5 + (g - 1) * 4 + k, col)] = f"{g}A"
    filled = {1: ["1 RAPIDI 1", "1 RAPIDI 2", "1 RAPIDI 3", "1 RAPIDI 4"],           # 1A dolu (rapidi)
              6: ["1 RAPIDI 1", "1 RAPIDI 2", "1 RAPIDI 3", "1 RAPIDI 4"],           # 6A dolu (rapidi)
              2: ["1 VİTRİFİTMAVİ", "1 RAPIDI 2", "1 RAPIDI 3"],                       # 2A: 1 boş satır, mavi kullanılmış
              4: ["1 CRYOLOCKSARI", "1 VİTRİFİTMAVİ"]}                                   # 4A: 2 boş satır
    for g, vials in filled.items():
        for k, v in enumerate(vials):
            put(1, 5 + (g - 1) * 4 + k, "SOYAD" + "ABCDEFGH"[g - 1], v)
    for g in range(1, 9):  # 2. canister: tamamen boş, ama soyad sütununda yeterli metin olsun diye 5-8. goblet dolu
        if g >= 5:
            for k in range(4):
                put(9, 5 + (g - 1) * 4 + k, "BSOYAD" + "ABCDEFGH"[g - 1], f"1 RAPIDI {k + 1}")
    return cells


def test_vial_helpers():
    assert vial_color("1 VİTRİFİTMAVİ") == "MAVİ" and vial_color("1 CRYOLOCKSARI") == "SARI"
    assert vial_color("1 vitrifityeşil") == "YEŞİL" and vial_color("1 CRYOTOPTURUNCU") == "TURUNCU"
    assert vial_color("1 RAPIDI 1") is None
    assert vial_text("CRYOLOCK", "SARI") == "1 CRYOLOCKSARI"


def test_scan_finds_positions_and_state():
    ps = scan_sheet("S", make_map_cells())
    assert len(ps) == 16 and all(len(p.rows) == 4 for p in ps)
    p = {(x.canister, x.label): x for x in ps}
    assert p[(1, "1A")].free_rows == [] and p[(1, "3A")].free_colors == COLORS
    assert len(p[(1, "2A")].free_rows) == 1 and p[(1, "2A")].free_colors == ["SARI", "YEŞİL", "TURUNCU"]
    assert p[(1, "4A")].used_colors == {"SARI", "MAVİ"} and p[(1, "2A")].tank == 1
    text = "\n".join(summarize(ps))
    assert "SOYAD" not in text and "Canister 1" in text


def test_suggest_two_straws_different_colors_adjacent():
    ps = [p for p in scan_sheet("S", make_map_cells()) if p.canister == 1]
    plan = suggest(ps, 2)[0]
    pl = plan.placements[0]
    assert plan.span == 1 and pl.position.label == "3A" and pl.colors == ["MAVİ", "SARI"]
    assert pl.rows[1].row == pl.rows[0].row + 1  # yan yana satırlar


def test_suggest_never_repeats_color_in_a_goblet_and_spans_neighbours():
    ps = [p for p in scan_sheet("S", make_map_cells()) if p.canister == 1]
    for n in (1, 2, 3, 5, 6):
        for plan in suggest(ps, n, limit=50):
            assert sum(len(pl.rows) for pl in plan.placements) == n
            for pl in plan.placements:
                taken = pl.position.used_colors
                assert len(set(pl.colors)) == len(pl.colors) and not (set(pl.colors) & taken)
                assert all(not r.occupied for r in pl.rows)
            nums = [pl.position.number for pl in plan.placements]
            assert nums == list(range(nums[0], nums[0] + len(nums)))  # komşu gobletler
    five = suggest(ps, 5)[0]
    assert five.span == 2  # 4'ten fazla straw tek goblete sığmaz


def test_suggest_colorless_uses_free_rows_only():
    ps = [p for p in scan_sheet("S", make_map_cells()) if p.canister == 1]
    plan = suggest(ps, 3, colored=False)[0]
    assert plan.placements[0].colors == [None, None, None]


def test_lower_levels_suggested_before_upper_levels():
    from tank_haritasi.harita import Position, Row
    def pos(suffix, can):
        label = f"1{suffix}"
        return Position("S", 1, can, label, 1, suffix,
                        [Row(i, {}, False, None, "") for i in range(1, 5)])
    ps = [pos("A", 1), pos("", 2)]            # canister 1'in ÜST katı, canister 2'nin ALT katı
    first = suggest(ps, 2)[0].placements[0].position
    assert first.kat == "alt" and first.canister == 2


def test_layout_detail_and_rapidi_not_unknown():
    from tank_haritasi.harita import describe_layout
    cells = make_map_cells()
    ps = scan_sheet("S", cells)
    text = "\n".join(describe_layout("S", cells, ps))
    assert "A1='TANK 1'" in text and "A3='CANISTER 1'" in text and "satır 5–36" in text and "SOYAD" not in text
    assert next(p for p in ps if p.canister == 1 and p.label == "1A").unknown_occupied == 0  # rapidi renksiz, belirsiz değil
