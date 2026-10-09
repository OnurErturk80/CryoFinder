"""Graph'ın kullandığımız bölümünü taklit eden bellek-içi sahte OneDrive."""
import io
import json
import re
import urllib.parse

from openpyxl import Workbook


def make_xlsx(cells: dict[str, object], sheet="Tank1") -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for k, v in cells.items():
        ws[k] = v
    b = io.BytesIO(); wb.save(b); return b.getvalue()


class Resp:
    def __init__(self, status=200, body=None, content=b""):
        self.status_code, self._body, self.content, self.headers = status, body, content, {}
        self.text = json.dumps(body) if body is not None else ""

    def json(self):
        return self._body


def err(status, code="x"):
    return Resp(status, {"error": {"code": code, "message": code}})


class FakeOneDrive:
    def __init__(self, files: dict[str, bytes], excel_api=False):
        self.items, self.excel_api, self.log, self.n = {}, excel_api, [], 0
        for path, data in files.items():
            self._put(path, data)

    def _put(self, path, data):
        self.n += 1
        iid = f"ID{self.n}"
        self.items[iid] = {"path": path, "data": data, "v": 1}
        return iid

    def _meta(self, iid):
        it = self.items[iid]
        m = {"id": iid, "name": it["path"].split("/")[-1], "eTag": f'"{iid},{it["v"]}"', "size": len(it["data"]), "file": {}}
        if it.get("folder"):
            m.pop("file"); m["folder"] = {}
        return m

    def by_path(self, p):
        return next((i for i, it in self.items.items() if it["path"] == p), None)

    def external_edit(self, path, data):  # başka biri dosyayı değiştirdi
        iid = self.by_path(path); self.items[iid]["data"] = data; self.items[iid]["v"] += 1

    def request(self, method, url, params=None, headers=None, json=None, data=None, timeout=0):
        self.log.append((method, url))
        headers = headers or {}
        u = urllib.parse.unquote(url.replace("https://graph.microsoft.com/v1.0", ""))
        if method == "GET" and (m := re.fullmatch(r"/me/drive/root:/(.+)", u)):
            iid = self.by_path(m[1]); return Resp(200, self._meta(iid)) if iid else err(404, "itemNotFound")
        if method == "GET" and (m := re.fullmatch(r"/me/drive/items/(\w+)", u)):
            return Resp(200, self._meta(m[1]))
        if method == "GET" and (m := re.fullmatch(r"/me/drive/items/(\w+)/content", u)):
            return Resp(200, content=self.items[m[1]]["data"])
        if method == "POST" and u == "/me/drive/root/children":
            self.n += 1; iid = f"ID{self.n}"
            self.items[iid] = {"path": json["name"], "data": b"", "v": 1, "folder": True}
            return Resp(201, self._meta(iid))
        if method == "PUT" and (m := re.fullmatch(r"/me/drive/root:/(.+):/content", u)):
            if self.by_path(m[1]):
                return err(409, "nameAlreadyExists")
            iid = self._put(m[1], data); return Resp(201, self._meta(iid))
        if method == "PUT" and (m := re.fullmatch(r"/me/drive/items/(\w+)/content", u)):
            it = self.items[m[1]]
            if headers.get("If-Match") != f'"{m[1]},{it["v"]}"':
                return err(412, "preconditionFailed")
            it["data"], it["v"] = data, it["v"] + 1
            return Resp(200, self._meta(m[1]))
        if "/workbook" in u:
            if not self.excel_api:
                return err(501, "notSupported")
            return self._excel(method, u, json)
        raise AssertionError(f"Sahte OneDrive bilmiyor: {method} {u}")

    def _excel(self, method, u, js):
        iid = re.search(r"items/(\w+)/workbook", u)[1]
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.items[iid]["data"]))
        if u.endswith("/worksheets"):
            return Resp(200, {"value": [{"id": f"{{{n}}}", "name": n} for n in wb.sheetnames]})
        if u.endswith("createSession"):
            return Resp(201, {"id": "S1"})
        if u.endswith("closeSession"):
            return Resp(204)
        m = re.search(r"worksheets/\{(.+?)\}/range\(address='(\w+)'\)", u)
        if m:
            ws, cell = wb[m[1]], m[2]
            if method == "PATCH":
                ws[cell] = js["values"][0][0]
                b = io.BytesIO(); wb.save(b)
                self.items[iid]["data"] = b.getvalue(); self.items[iid]["v"] += 1
                return Resp(200, {})
            v = ws[cell].value
            return Resp(200, {"values": [[v if v is not None else ""]], "formulas": [[v if v is not None else ""]]})
        m = re.search(r"worksheets/\{(.+?)\}/usedRange", u)
        if m:
            ws = wb[m[1]]
            return Resp(200, {"address": f"{m[1]}!A1:{ws.dimensions.split(':')[-1]}",
                              "values": [[c.value if c.value is not None else "" for c in row] for row in ws.iter_rows()]})
        raise AssertionError(u)
